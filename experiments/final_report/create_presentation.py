"""Creates the AFM-RL Final Presentation PowerPoint (AFM_RL_Final_Presentation.pptx).

Generates a modern, clean, 16-slide PowerPoint presentation with:
- 16:9 widescreen layout
- Professional navy/blue/slate color palette
- High-contrast text cards and callouts
- Embedded high-resolution comparison plots
- Verified final experimental results
"""

from pathlib import Path
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

REPORT_DIR = Path(__file__).resolve().parent
PLOTS_DIR = REPORT_DIR / "plots"
OUTPUT_PPTX = REPORT_DIR / "AFM_RL_Final_Presentation.pptx"

# -----------------------------------------------------------------------------
# Design Theme & Color Palette
# -----------------------------------------------------------------------------
BG_COLOR       = RGBColor(248, 250, 252)   # Slate-50 (Very light gray-blue)
NAVY_PRIMARY   = RGBColor(15, 23, 42)      # Slate-900 (Dark navy)
BLUE_ACCENT    = RGBColor(30, 58, 138)     # Blue-900 (Deep royal blue)
BLUE_HIGHLIGHT = RGBColor(37, 99, 235)     # Blue-600 (Vibrant blue)
TEXT_MAIN      = RGBColor(30, 41, 59)      # Slate-800
TEXT_MUTED     = RGBColor(100, 116, 139)   # Slate-500
CARD_BG        = RGBColor(255, 255, 255)   # White
CARD_BORDER    = RGBColor(226, 232, 240)   # Slate-200
GREEN_ACCENT   = RGBColor(22, 101, 52)     # Emerald-800
AMBER_ACCENT   = RGBColor(146, 64, 14)     # Amber-800


def create_deck():
    prs = Presentation()
    # 16:9 Widescreen (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_blank_slide(title_text="", category_text="AFM-RL: ADAPTIVE FEATURE MATCHING"):
        slide = prs.slides.add_slide(blank_layout)
        
        # Background fill
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()

        # Category eyebrow
        if category_text:
            cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
            tf_c = cat_box.text_frame
            tf_c.word_wrap = True
            p_c = tf_c.paragraphs[0]
            p_c.text = category_text.upper()
            p_c.font.size = Pt(10.5)
            p_c.font.bold = True
            p_c.font.color.rgb = BLUE_HIGHLIGHT

        # Title header
        if title_text:
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
            tf_t = title_box.text_frame
            tf_t.word_wrap = True
            p_t = tf_t.paragraphs[0]
            p_t.text = title_text
            p_t.font.size = Pt(22)
            p_t.font.bold = True
            p_t.font.color.rgb = NAVY_PRIMARY

        return slide

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        return card

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY_PRIMARY
    bg1.line.fill.background()

    # Title box
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "Adaptive Feature Matching using Hierarchical Reinforcement Learning"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.space_after = Pt(16)

    p2 = tf1.add_paragraph()
    p2.text = "Level-2 Adaptive Keypoint Budget Optimization & Benchmark Evaluation"
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(147, 197, 253)
    p2.space_after = Pt(24)

    p3 = tf1.add_paragraph()
    p3.text = "Final Project Report | Verified Results across HPatches & MegaDepth-1500"
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(203, 213, 225)

    # =========================================================================
    # SLIDE 2: The Problem
    # =========================================================================
    s2 = add_blank_slide("The Problem: How Many Keypoints Should We Use?", "MOTIVATION & CHALLENGE")
    
    # Left card
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb_l = s2.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tfl = tb_l.text_frame
    tfl.word_wrap = True
    
    p = tfl.paragraphs[0]
    p.text = "The Keypoint Dilemma"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(14)
    
    bullets_l = [
        ("Traditional pipelines use a fixed number of keypoints", "Every image pair gets the exact same budget (e.g., 1000 or 4000 keypoints)."),
        ("Too few keypoints (e.g. 500):", "Hard pairs fail to match when lighting or perspective changes significantly."),
        ("Too many keypoints (e.g. 4000):", "Easy pairs waste massive compute and extract redundant/noisy features."),
        ("Core Reality:", "No single fixed budget is optimal across diverse real-world images.")
    ]
    for b_title, b_desc in bullets_l:
        p = tfl.add_paragraph()
        p.text = f"• {b_title}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = NAVY_PRIMARY
        p_sub = tfl.add_paragraph()
        p_sub.text = f"   {b_desc}"
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.space_after = Pt(8)

    # Right card
    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), bg_color=RGBColor(241, 245, 249))
    tb_r = s2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tfr = tb_r.text_frame
    tfr.word_wrap = True
    
    p = tfr.paragraphs[0]
    p.text = "What We Need"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT
    p.space_after = Pt(14)
    
    needs = [
        "Dynamic allocation: High budget for complex viewpoint changes.",
        "Early stopping: Low budget for simple illumination variations.",
        "Zero manual tuning: An agent that automatically senses difficulty.",
        "Compute efficiency: Maximize accuracy while minimizing feature extraction cost."
    ]
    for n in needs:
        p = tfr.add_paragraph()
        p.text = f"✓  {n}"
        p.font.size = Pt(13.5)
        p.font.color.rgb = NAVY_PRIMARY
        p.space_after = Pt(14)

    # =========================================================================
    # SLIDE 3: Our Idea
    # =========================================================================
    s3 = add_blank_slide("Our Solution: Adaptive Keypoint Budget via Reinforcement Learning", "CORE ARCHITECTURE")
    
    card_w = Inches(3.6)
    cards_data = [
        ("Level 1: Pipeline Selection", "Selects the best detector and descriptor algorithm (e.g., SIFT, AKAZE, ORB) suited for the overall scene.", "(Future Integration / Modular)"),
        ("Level 2: Budget Allocation", "Dynamically adjusts the keypoint count (+500, +1000, -500, STOP) for the chosen pipeline.", "(Focus of this Evaluation)"),
        ("Utility-Driven Objective", "Maximizes matching accuracy while imposing a keypoint cost penalty (λ = 0.1).", "(Learns When to Stop)")
    ]
    for i, (head, body, sub) in enumerate(cards_data):
        c_left = Inches(0.8 + i * 4.0)
        add_card(s3, c_left, Inches(1.8), card_w, Inches(4.8))
        tb = s3.shapes.add_textbox(c_left + Inches(0.2), Inches(2.1), card_w - Inches(0.4), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"0{i+1}"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = BLUE_HIGHLIGHT
        p.space_after = Pt(8)
        
        p = tf.add_paragraph()
        p.text = head
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p.space_after = Pt(12)
        
        p = tf.add_paragraph()
        p.text = body
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(14)
        
        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = BLUE_ACCENT

    # =========================================================================
    # SLIDE 4: How the Level-2 Agent Works
    # =========================================================================
    s4 = add_blank_slide("Inside the Level-2 RL Agent: Step-by-Step Execution", "RL AGENT WORKFLOW")
    
    steps_data = [
        ("1. Observe", "Reads 31-D state vector:\n• Image pair statistics (contrast, brightness)\n• Pipeline 1-hot encoding\n• Matching score, inlier ratio, recent delta"),
        ("2. Action", "PPO Policy chooses 1 of 4 discrete actions:\n• +500 keypoints\n• +1,000 keypoints\n• -500 keypoints\n• STOP immediately"),
        ("3. Evaluate", "Runs detector/descriptor pool slice, executes BF match, and fits homography via RANSAC."),
        ("4. Reward", "Reward = Δ Utility - Step Penalty\nUtility = Accuracy - λ × (Keypoints / 4000)\nAgent learns to stop when gains saturate.")
    ]
    for i, (title, desc) in enumerate(steps_data):
        c_left = Inches(0.8 + i * 2.95)
        add_card(s4, c_left, Inches(1.8), Inches(2.8), Inches(4.9))
        tb = s4.shapes.add_textbox(c_left + Inches(0.15), Inches(2.0), Inches(2.5), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = BLUE_HIGHLIGHT
        p.space_after = Pt(10)
        
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 5: Why Fixed Budgets Are Not Ideal
    # =========================================================================
    s5 = add_blank_slide("Why Fixed Budgets Fail: The Baseline Comparison", "BASELINE MOTIVATION")
    
    budgets = [
        ("Fixed 500", "495.6 KP Avg", "26.89% AUC@3px", "Lightweight but suffers from low inlier count on viewpoint changes.", False),
        ("Fixed 1000", "979.5 KP Avg", "31.08% AUC@3px", "Standard default, but still fails on challenging perspective shifts.", False),
        ("Fixed 2000", "1,897.0 KP Avg", "35.07% AUC@3px", "Higher accuracy, but wastes compute on simple illumination pairs.", False),
        ("Fixed 4000", "3,424.4 KP Avg", "36.34% AUC@3px", "Brute-force upper bound; very slow and extracts redundant points.", False),
    ]
    for i, (name, kp, auc, desc, highlight) in enumerate(budgets):
        c_left = Inches(0.8 + i * 2.95)
        add_card(s5, c_left, Inches(1.8), Inches(2.8), Inches(4.9))
        tb = s5.shapes.add_textbox(c_left + Inches(0.15), Inches(2.0), Inches(2.5), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = name
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        
        p = tf.add_paragraph()
        p.text = auc
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = BLUE_HIGHLIGHT
        p.space_after = Pt(6)
        
        p = tf.add_paragraph()
        p.text = f"Keypoints: {kp}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(10)
        
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 6: Experimental Setup
    # =========================================================================
    s6 = add_blank_slide("Experimental Setup: HPatches & RIPE Evaluation Protocol", "BENCHMARK DESIGN")
    
    # Left Card: Dataset info
    add_card(s6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb_l = s6.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tfl = tb_l.text_frame
    tfl.word_wrap = True
    
    p = tfl.paragraphs[0]
    p.text = "HPatches Benchmark Split"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(12)
    
    items = [
        "116 Total Sequences in HPatches dataset.",
        "Held-out Test Split: 17 sequences = 85 image pairs.",
        "8 Illumination sequences (40 pairs): Brightness variations.",
        "9 Viewpoint sequences (45 pairs): Wide camera angles.",
        "Strict Isolation: Zero train/validation overlap with test set.",
        "4,675 Total Evaluations: 85 pairs × 11 presets × 5 methods."
    ]
    for it in items:
        p = tfl.add_paragraph()
        p.text = f"• {it}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(6)

    # Right Card: Metric info
    add_card(s6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb_r = s6.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tfr = tb_r.text_frame
    tfr.word_wrap = True
    
    p = tfr.paragraphs[0]
    p.text = "RIPE Homography AUC Metric"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(12)
    
    m_items = [
        "Mean 4-Corner Reprojection Error: Pixel distance between estimated homography corners and ground truth.",
        "Normalized Cumulative AUC (Trapezoidal Integration):",
        "  • AUC@1px: High precision registration (< 1 px).",
        "  • AUC@3px: Primary benchmark threshold (< 3 px).",
        "  • AUC@5px: General geometric alignment (< 5 px).",
        "Success Rates: Percentage of pairs with error < threshold."
    ]
    for mit in m_items:
        p = tfr.add_paragraph()
        p.text = f"• {mit}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 7: Main Result — Overall AUC@3
    # =========================================================================
    s7 = add_blank_slide("Main Result: Adaptive Level-2 Outperforms All Fixed Baselines", "HPATCHES OVERALL PERFORMANCE")
    
    # Image on left
    img_path = str(PLOTS_DIR / "hpatches_auc3_overall.png")
    if Path(img_path).exists():
        s7.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), width=Inches(6.4))
    
    # Summary card on right
    add_card(s7, Inches(7.5), Inches(1.6), Inches(5.0), Inches(5.2))
    tb = s7.shapes.add_textbox(Inches(7.8), Inches(1.9), Inches(4.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Key Verified Numbers"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(12)
    
    res_items = [
        ("Adaptive Level-2:", "38.30% AUC@3px (1,192 KP avg)"),
        ("Fixed 4000:", "36.34% AUC@3px (3,424 KP avg)"),
        ("Fixed 2000:", "35.07% AUC@3px (1,897 KP avg)"),
        ("Fixed 1000:", "31.08% AUC@3px (980 KP avg)"),
        ("Fixed 500:", "26.89% AUC@3px (496 KP avg)"),
    ]
    for k, v in res_items:
        p = tf.add_paragraph()
        p.text = f"{k} "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = NAVY_PRIMARY
        p_v = tf.add_paragraph()
        p_v.text = f"  {v}"
        p_v.font.size = Pt(12.5)
        p_v.font.color.rgb = BLUE_HIGHLIGHT
        p_v.space_after = Pt(6)
        
    p_sum = tf.add_paragraph()
    p_sum.text = "Outcome: Adaptive Level-2 achieves highest overall accuracy while using ~65% fewer keypoints than Fixed 4000."
    p_sum.font.size = Pt(12)
    p_sum.font.bold = True
    p_sum.font.color.rgb = GREEN_ACCENT

    # =========================================================================
    # SLIDE 8: Main Result — Accuracy vs Keypoint Cost
    # =========================================================================
    s8 = add_blank_slide("Accuracy vs Keypoint Cost: 65% Computational Reduction", "BUDGET EFFICIENCY")
    
    img_path = str(PLOTS_DIR / "hpatches_keypoints_overall.png")
    if Path(img_path).exists():
        s8.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), width=Inches(6.4))
        
    add_card(s8, Inches(7.5), Inches(1.6), Inches(5.0), Inches(5.2))
    tb = s8.shapes.add_textbox(Inches(7.8), Inches(1.9), Inches(4.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Efficiency Breakdown"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(12)
    
    eff_items = [
        ("vs Fixed 4000:", "Saves 2,233 keypoints per pair on average (-65% budget reduction) while increasing AUC@3px by +1.96%."),
        ("vs Fixed 2000:", "Saves 705 keypoints per pair on average (-37% budget reduction) while increasing AUC@3px by +3.23%."),
        ("vs Fixed 500:", "Gains +11.41% AUC@3px and +21.72% in success rate by dynamically scaling up only when required."),
        ("Adaptation by Difficulty:", "Illumination pairs average 934 KP; Viewpoint pairs scale to 1,421 KP.")
    ]
    for k, v in eff_items:
        p = tf.add_paragraph()
        p.text = f"• {k} "
        p.font.bold = True
        p.font.size = Pt(12.5)
        p.font.color.rgb = NAVY_PRIMARY
        p_v = tf.add_paragraph()
        p_v.text = f"  {v}"
        p_v.font.size = Pt(11.5)
        p_v.font.color.rgb = TEXT_MAIN
        p_v.space_after = Pt(6)

    # =========================================================================
    # SLIDE 9: Performance Across All 11 Presets
    # =========================================================================
    s9 = add_blank_slide("Performance Across All 11 Feature Matching Presets", "PRESET-LEVEL ANALYSIS")
    
    img_path = str(PLOTS_DIR / "hpatches_auc3_by_preset.png")
    if Path(img_path).exists():
        s9.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), width=Inches(11.7))

    # =========================================================================
    # SLIDE 10: Best and Worst Presets
    # =========================================================================
    s10 = add_blank_slide("Preset Extremes: SIFT (Best) vs ORB+FREAK (Worst)", "PIPELINE COMPARISON")
    
    # Left Card: SIFT
    add_card(s10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), border_color=RGBColor(187, 247, 208))
    tb_l = s10.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tfl = tb_l.text_frame
    tfl.word_wrap = True
    
    p = tfl.paragraphs[0]
    p.text = "🏆 Best Preset: SIFT"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT
    p.space_after = Pt(10)
    
    sift_bullets = [
        "AUC@1px: 17.44% | AUC@3px: 50.03% | AUC@5px: 64.41%",
        "Success Rate @ 3px: 78.82% of all test pairs.",
        "Average Keypoints: 942 KP (2.94 steps).",
        "Why it excelled: SIFT's 128-D float descriptors provide high distinctive matching. The agent learns that SIFT achieves high inlier ratio quickly and terminates early, saving budget."
    ]
    for b in sift_bullets:
        p = tfl.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    # Right Card: ORB+FREAK
    add_card(s10, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), border_color=RGBColor(254, 202, 202))
    tb_r = s10.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tfr = tb_r.text_frame
    tfr.word_wrap = True
    
    p = tfr.paragraphs[0]
    p.text = "⚠️ Lowest Preset: ORB+FREAK"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = AMBER_ACCENT
    p.space_after = Pt(10)
    
    orb_bullets = [
        "AUC@1px: 3.73% | AUC@3px: 21.37% | AUC@5px: 33.69%",
        "Success Rate @ 3px: 44.71% of all test pairs.",
        "Average Keypoints: 1,235 KP (4.01 steps).",
        "Measured finding: ORB+FREAK produced the lowest measured AUC values on this benchmark. This tells us this specific hybrid pipeline was more difficult for the system on this dataset."
    ]
    for b in orb_bullets:
        p = tfr.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    # =========================================================================
    # SLIDE 11: Why Adaptive is Useful
    # =========================================================================
    s11 = add_blank_slide("Why Adaptive Level-2 is Superior: Head-to-Head vs Fixed-4000", "SUMMARY OF GAINS")
    
    img_path = str(PLOTS_DIR / "hpatches_adaptive_vs_fixed4000.png")
    if Path(img_path).exists():
        s11.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), width=Inches(7.0))
        
    add_card(s11, Inches(8.1), Inches(1.6), Inches(4.4), Inches(5.2))
    tb = s11.shapes.add_textbox(Inches(8.3), Inches(1.8), Inches(4.0), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "10 of 11 Presets Match or Beat Fixed 4000"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(12)
    
    wins = [
        "AKAZE: 42.71% vs 36.69% (+6.02%)",
        "BRISK: 42.43% vs 38.96% (+3.47%)",
        "GFTT+SIFT: 44.61% vs 41.49% (+3.12%)",
        "GFTT+BRIEF: 36.09% vs 32.98% (+3.11%)",
        "FAST+BRIEF: 36.01% vs 33.60% (+2.41%)",
        "KAZE: 45.49% vs 44.30% (+1.19%)",
        "ORB: 33.72% vs 33.34% (+0.38%)",
        "STAR+BRIEF: 30.21% vs 29.83% (+0.38%)",
        "ORB+FREAK: 21.37% vs 17.69% (+3.68%)",
        "FAST+FREAK: 38.64% vs 37.99% (+0.65%)",
    ]
    for w in wins[:8]:
        p = tf.add_paragraph()
        p.text = f"• {w}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 12: MegaDepth Experiment
    # =========================================================================
    s12 = add_blank_slide("MegaDepth-1500 Experiment: Testing 3D Relative Pose Estimation", "OUTDOOR 3D BENCHMARK")
    
    add_card(s12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb_l = s12.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tfl = tb_l.text_frame
    tfl.word_wrap = True
    
    p = tfl.paragraphs[0]
    p.text = "Why MegaDepth-1500?"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(12)
    
    md_why = [
        "HPatches evaluates planar homographies (2D transformations).",
        "MegaDepth evaluates 3D outdoor scenes with 6-DoF camera motion.",
        "Goal: Test whether the Level-2 RL agent can be trained directly on epipolar geometry.",
        "Dataset Split: 1,050 train pairs, 225 validation pairs, 225 test pairs.",
        "Training Budget: 170,528 timesteps matching the verified HPatches budget."
    ]
    for it in md_why:
        p = tfl.add_paragraph()
        p.text = f"• {it}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(6)

    add_card(s12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb_r = s12.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tfr = tb_r.text_frame
    tfr.word_wrap = True
    
    p = tfr.paragraphs[0]
    p.text = "3D Relative Pose Metric"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(12)
    
    md_met = [
        "Essential Matrix Estimation via 5-point RANSAC using camera intrinsics (K1, K2).",
        "Angular Pose Error: Maximum of rotation error and translation angle error: max(err_R, err_t).",
        "Pose AUC Reported at 5°, 10°, and 20° thresholds (RIPE protocol standard).",
        "Training Reward: Sampson epipolar distance surrogate: exp(-mean_sampson_error / 5.0)."
    ]
    for it in md_met:
        p = tfr.add_paragraph()
        p.text = f"• {it}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 13: MegaDepth Result
    # =========================================================================
    s13 = add_blank_slide("MegaDepth Results: Honest Analysis of Surrogate Reward", "MEGADEPTH TEST RESULTS")
    
    img_path = str(PLOTS_DIR / "megadepth_auc20_comparison.png")
    if Path(img_path).exists():
        s13.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), width=Inches(6.4))
        
    add_card(s13, Inches(7.5), Inches(1.6), Inches(5.0), Inches(5.2))
    tb = s13.shapes.add_textbox(Inches(7.8), Inches(1.9), Inches(4.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Key Findings on MegaDepth"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(12)
    
    md_res = [
        ("MegaDepth-Trained Adaptive:", "26.02% AUC@20° (529 KP avg)"),
        ("HPatches Zero-Shot Adaptive:", "35.39% AUC@20° (619 KP avg)"),
        ("Fixed 500 Baseline:", "34.13% AUC@20°"),
        ("Fixed 4000 Baseline:", "62.08% AUC@20°"),
    ]
    for k, v in md_res:
        p = tf.add_paragraph()
        p.text = f"{k} "
        p.font.bold = True
        p.font.size = Pt(12.5)
        p.font.color.rgb = NAVY_PRIMARY
        p_v = tf.add_paragraph()
        p_v.text = f"  {v}"
        p_v.font.size = Pt(12)
        p_v.font.color.rgb = TEXT_MUTED
        p_v.space_after = Pt(4)
        
    p_exp = tf.add_paragraph()
    p_exp.text = "Interpretation: The simple Sampson-distance surrogate reward was insufficient to guide productive policy learning. The policy became overly conservative."
    p_exp.font.size = Pt(11.5)
    p_exp.font.bold = True
    p_exp.font.color.rgb = AMBER_ACCENT

    # =========================================================================
    # SLIDE 14: What We Learned
    # =========================================================================
    s14 = add_blank_slide("Key Insights: What Did We Learn?", "LESSONS & INSIGHTS")
    
    insights = [
        ("1. Dynamic keypoint allocation is highly effective for 2D homography", "Achieved higher accuracy than 4,000 fixed points while saving 65% of keypoint budget."),
        ("2. SIFT remains the strongest classical matching pipeline", "Yielded 50.03% AUC@3px on HPatches test split; stopping early when inliers saturate."),
        ("3. Fixed budgets are fundamentally inefficient", "Fixed 500 fails on hard scenes; Fixed 4000 is heavily redundant on clean scenes."),
        ("4. Reward formulation is critical for RL transfer", "Using surrogate distance rewards for 3D pose estimation led to conservative budget stopping. Future 3D models require direct Essential matrix pose rewards.")
    ]
    for i, (title, desc) in enumerate(insights):
        top_pos = Inches(1.7 + i * 1.3)
        add_card(s14, Inches(0.8), top_pos, Inches(11.7), Inches(1.15))
        tb = s14.shapes.add_textbox(Inches(1.1), top_pos + Inches(0.1), Inches(11.1), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 15: Limitations and Future Work
    # =========================================================================
    s15 = add_blank_slide("Project Scope: Limitations & Future Directions", "ROADMAP")
    
    # Left Card: Limitations
    add_card(s15, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb_l = s15.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tfl = tb_l.text_frame
    tfl.word_wrap = True
    
    p = tfl.paragraphs[0]
    p.text = "Current Limitations"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(12)
    
    lims = [
        "Level 2 Only: Focus was on budget allocation; Level 1 pipeline selection was not trained end-to-end.",
        "MegaDepth Reward Surrogate: Utilized Sampson distance rather than exact 5-point angular error.",
        "Hardware Profiling: Efficiency measured via keypoint count, not physical millisecond runtime on edge chips.",
        "Aachen Out of Scope: 6-DoF visual localization benchmark was not evaluated in this phase."
    ]
    for it in lims:
        p = tfl.add_paragraph()
        p.text = f"• {it}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    # Right Card: Future Work
    add_card(s15, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb_r = s15.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tfr = tb_r.text_frame
    tfr.word_wrap = True
    
    p = tfr.paragraphs[0]
    p.text = "Future Work"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT
    p.space_after = Pt(12)
    
    futs = [
        "Level 1 + Level 2 Integration: Train high-level pipeline selection on top of frozen Level 2 budget policy.",
        "Direct 3D Pose Rewards: Incorporate camera intrinsics into MegaDepth step rewards.",
        "Aachen Day-Night Evaluation: Benchmark on outdoor visual localization under extreme illumination shifts.",
        "Edge Device Deployment: Measure wall-clock latency on NVIDIA Jetson and mobile devices."
    ]
    for it in futs:
        p = tfr.add_paragraph()
        p.text = f"• {it}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    # =========================================================================
    # SLIDE 16: Final Conclusion
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    bg16 = s16.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg16.fill.solid()
    bg16.fill.fore_color.rgb = NAVY_PRIMARY
    bg16.line.fill.background()

    tb16 = s16.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(10.9), Inches(4.5))
    tf16 = tb16.text_frame
    tf16.word_wrap = True
    
    p = tf16.paragraphs[0]
    p.text = "CONCLUSION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(147, 197, 253)
    p.space_after = Pt(14)
    
    p = tf16.add_paragraph()
    p.text = "Adaptive Reinforcement Learning Solves the Fixed Keypoint Dilemma"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.space_after = Pt(20)
    
    p = tf16.add_paragraph()
    p.text = (
        "On the HPatches benchmark under the RIPE evaluation protocol, our Level-2 Adaptive RL agent achieved "
        "38.30% AUC@3px, beating the brute-force Fixed-4000 baseline (36.34%) while saving over 65% of keypoint "
        "extractions (1,192 vs 3,424 keypoints).\n\n"
        "By sensing scene difficulty dynamically, the agent scales up keypoints on challenging viewpoint changes "
        "and stops early on clean illumination pairs, providing superior accuracy and computational efficiency."
    )
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(226, 232, 240)

    prs.save(str(OUTPUT_PPTX))
    print(f"Saved {OUTPUT_PPTX}")


if __name__ == "__main__":
    create_deck()
